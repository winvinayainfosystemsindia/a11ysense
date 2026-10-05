"""
Accessibility Audit Quotation Generator for A11ySense AI.
Produces executive Excel (.xlsx) and PDF (.pdf) quotation proposals.
"""
import io
import math
from datetime import datetime, timedelta
from typing import Dict, Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ReportLab imports for PDF generation
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

# ── Excel Colors & Styles ──────────────────────────────────────────────────
NAVY_HEADER_FILL = PatternFill(start_color="0D1B2A", end_color="0D1B2A", fill_type="solid")
PRIMARY_BLUE_FILL = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
ACCENT_BLUE_FILL = PatternFill(start_color="1565C0", end_color="1565C0", fill_type="solid")
TOTAL_ROW_FILL = PatternFill(start_color="E8EAF6", end_color="E8EAF6", fill_type="solid")
FINAL_QUOTE_FILL = PatternFill(start_color="1B5E20", end_color="1B5E20", fill_type="solid")

CARD_BG_FILL = PatternFill(start_color="F4F6F9", end_color="F4F6F9", fill_type="solid")
CARD_HEADER_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

SIMPLE_FILL = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
SIMPLE_FONT = Font(name="Calibri", size=10, color="2E7D32", bold=True)

MEDIUM_FILL = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
MEDIUM_FONT = Font(name="Calibri", size=10, color="E65100", bold=True)

COMPLEX_FILL = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
COMPLEX_FONT = Font(name="Calibri", size=10, color="C62828", bold=True)

WHITE_BOLD_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
WHITE_TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=15, bold=True, color="0D1B2A")
SUBTITLE_FONT = Font(name="Calibri", size=10, color="555555")
SECTION_FONT = Font(name="Calibri", size=12, bold=True, color="1A237E")
REGULAR_FONT = Font(name="Calibri", size=10)
BOLD_FONT = Font(name="Calibri", size=10, bold=True)
CODE_FONT = Font(name="Consolas", size=9)
LINK_FONT = Font(name="Calibri", size=10, color="0563C1", underline="single")

THIN_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
HEADER_BORDER = Border(
    left=Side(style="thin", color="94A3B8"),
    right=Side(style="thin", color="94A3B8"),
    top=Side(style="thin", color="94A3B8"),
    bottom=Side(style="medium", color="0D1B2A"),
)
TOTAL_DOUBLE_BORDER = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="double", color="1A237E"),
)


def _safe_val(val: Any) -> Any:
    if val is None:
        return ""
    if isinstance(val, (int, float, bool)):
        return val
    s = str(val)
    if s and s[0] in ('=', '+', '-', '@'):
        return "'" + s
    return s


def generate_quotation_excel(data: Dict[str, Any]) -> io.BytesIO:
    """
    Generates a commercial 3-sheet Excel Quotation Proposal:
    Sheet 1: Executive Quotation & Commercial Proposal
    Sheet 2: Page Line-Items & Effort Breakdown
    Sheet 3: Technical Element Evidence
    """
    wb = openpyxl.Workbook()
    # Sheet 1: Quotation Proposal
    ws1 = wb.active
    ws1.title = "Commercial Proposal"
    ws1.views.sheetView[0].showGridLines = True

    # Sheet 2: Page Breakdown
    ws2 = wb.create_sheet(title="Page Line-Items")
    ws2.views.sheetView[0].showGridLines = True

    # Sheet 3: Element Inventory
    ws3 = wb.create_sheet(title="Element Inventory")
    ws3.views.sheetView[0].showGridLines = True

    summary = data.get("summary", {})
    pages = data.get("pages", [])
    now = datetime.now()
    quote_ref = f"A11Y-QUOTE-{now.strftime('%Y%m%d')}-{abs(hash(str(summary))) % 10000:04d}"
    valid_until = (now + timedelta(days=30)).strftime("%B %d, %Y")

    # ══════════════════════════════════════════════════════════════════════════
    # SHEET 1: COMMERCIAL PROPOSAL
    # ══════════════════════════════════════════════════════════════════════════
    ws1.column_dimensions['A'].width = 4
    ws1.column_dimensions['B'].width = 38
    ws1.column_dimensions['C'].width = 22
    ws1.column_dimensions['D'].width = 18
    ws1.column_dimensions['E'].width = 20
    ws1.column_dimensions['F'].width = 22
    ws1.column_dimensions['G'].width = 4

    # Top Header Banner
    ws1.row_dimensions[2].height = 42
    ws1.merge_cells('B2:F2')
    top_cell = ws1['B2']
    top_cell.value = "  A11ySense AI — Accessibility Audit Commercial Quotation"
    top_cell.fill = NAVY_HEADER_FILL
    top_cell.font = WHITE_TITLE_FONT
    top_cell.alignment = Alignment(vertical="center", horizontal="left")

    # Subtitle / Metadata
    target_url = pages[0].get("url", "") if pages else ""
    metadata_rows = [
        ("Quotation Reference:", quote_ref, "Date Issued:", now.strftime("%B %d, %Y")),
        ("Target Website:", target_url, "Proposal Validity:", valid_until),
        ("Standard / Compliance:", "WCAG 2.2 Level A & AA / Section 508 / ADA", "Pricing Model:", "Fixed-Scope Commercial Quote"),
    ]

    cur_row = 4
    for label1, val1, label2, val2 in metadata_rows:
        ws1.row_dimensions[cur_row].height = 20
        c_l1 = ws1.cell(row=cur_row, column=2, value=label1)
        c_l1.font = BOLD_FONT
        c_l1.alignment = Alignment(vertical="center")

        c_v1 = ws1.cell(row=cur_row, column=3, value=val1)
        c_v1.font = REGULAR_FONT
        c_v1.alignment = Alignment(vertical="center")

        c_l2 = ws1.cell(row=cur_row, column=4, value=label2)
        c_l2.font = BOLD_FONT
        c_l2.alignment = Alignment(vertical="center")

        c_v2 = ws1.cell(row=cur_row, column=5, value=val2)
        c_v2.font = REGULAR_FONT
        c_v2.alignment = Alignment(vertical="center")
        cur_row += 1

    cur_row += 1

    # 4 KPI Summary Cards
    ws1.row_dimensions[cur_row].height = 18
    ws1.row_dimensions[cur_row + 1].height = 32

    # Card 1: Final Quoted Investment
    ws1.merge_cells(f'B{cur_row}:B{cur_row}')
    ws1.cell(row=cur_row, column=2, value="TOTAL QUOTED INVESTMENT").fill = CARD_HEADER_FILL
    ws1.cell(row=cur_row, column=2).font = Font(name="Calibri", size=9, bold=True, color="555555")
    ws1.cell(row=cur_row, column=2).alignment = Alignment(horizontal="center", vertical="center")
    c_k1 = ws1.cell(row=cur_row + 1, column=2, value=summary.get("total_final_cost", 0.0))
    c_k1.font = Font(name="Calibri", size=18, bold=True, color="1B5E20")
    c_k1.number_format = "$#,##0.00"
    c_k1.fill = CARD_BG_FILL
    c_k1.alignment = Alignment(horizontal="center", vertical="center")

    # Card 2: Manual Audit Effort
    ws1.cell(row=cur_row, column=3, value="MANUAL AUDIT EFFORT").fill = CARD_HEADER_FILL
    ws1.cell(row=cur_row, column=3).font = Font(name="Calibri", size=9, bold=True, color="555555")
    ws1.cell(row=cur_row, column=3).alignment = Alignment(horizontal="center", vertical="center")
    c_k2 = ws1.cell(row=cur_row + 1, column=3, value=f"{summary.get('total_manual_hours', 0.0):.1f} hrs")
    c_k2.font = Font(name="Calibri", size=18, bold=True, color="1A237E")
    c_k2.fill = CARD_BG_FILL
    c_k2.alignment = Alignment(horizontal="center", vertical="center")

    # Card 3: Scope Breakdown
    ws1.cell(row=cur_row, column=4, value="AUDIT SCOPE").fill = CARD_HEADER_FILL
    ws1.cell(row=cur_row, column=4).font = Font(name="Calibri", size=9, bold=True, color="555555")
    ws1.cell(row=cur_row, column=4).alignment = Alignment(horizontal="center", vertical="center")
    c_k3 = ws1.cell(row=cur_row + 1, column=4, value=f"{summary.get('total_pages', 0)} Page(s)")
    c_k3.font = Font(name="Calibri", size=18, bold=True, color="333333")
    c_k3.fill = CARD_BG_FILL
    c_k3.alignment = Alignment(horizontal="center", vertical="center")

    # Card 4: Base Cost (Before Margin)
    ws1.cell(row=cur_row, column=5, value="BASE SERVICE COST").fill = CARD_HEADER_FILL
    ws1.cell(row=cur_row, column=5).font = Font(name="Calibri", size=9, bold=True, color="555555")
    ws1.cell(row=cur_row, column=5).alignment = Alignment(horizontal="center", vertical="center")
    c_k4 = ws1.cell(row=cur_row + 1, column=5, value=summary.get("total_base_cost", 0.0))
    c_k4.font = Font(name="Calibri", size=18, bold=True, color="555555")
    c_k4.number_format = "$#,##0.00"
    c_k4.fill = CARD_BG_FILL
    c_k4.alignment = Alignment(horizontal="center", vertical="center")

    cur_row += 3

    # Section Header: Commercial Breakdown Table
    ws1.cell(row=cur_row, column=2, value="COMMERCIAL PROPOSAL BREAKDOWN").font = SECTION_FONT
    cur_row += 1

    table_headers = ["Deliverable / Service Component", "Scope / Quantity", "Unit Rate", "Base Subtotal", "Quoted Total (+Margin)"]
    ws1.row_dimensions[cur_row].height = 26
    for c_idx, h_text in enumerate(table_headers, start=2):
        cell = ws1.cell(row=cur_row, column=c_idx, value=h_text)
        cell.fill = PRIMARY_BLUE_FILL
        cell.font = WHITE_BOLD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = HEADER_BORDER

    cur_row += 1

    # Line Items
    margin_pct = summary.get("profit_margin_pct", 30.0)
    margin_factor = 1.0 + (margin_pct / 100.0)

    manual_hrs = summary.get("total_manual_hours", 0.0)
    hourly_rate = summary.get("hourly_rate", 40.0)
    manual_base = summary.get("total_manual_cost", 0.0)
    manual_quoted = manual_base * margin_factor

    api_base = summary.get("total_api_cost", 0.0)
    api_quoted = api_base * margin_factor

    plat_base = summary.get("total_platform_cost", 0.0)
    plat_quoted = plat_base * margin_factor

    total_pages = summary.get("total_pages", len(pages))

    line_items = [
        ("Manual Specialist Verification (NVDA / JAWS / Keyboard Audit)", f"{manual_hrs:.1f} Hours", f"${hourly_rate:.2f} / hr", manual_base, manual_quoted),
        ("Automated WCAG 2.2 Model & Rules Processing", f"{total_pages} Page(s)", "Tiered by complexity", api_base, api_quoted),
        ("Cloud Headless Browser Orchestration & Compute", f"{total_pages} Page(s)", f"${summary.get('platform_fee_per_page', 10.0):.2f} / page", plat_base, plat_quoted),
    ]

    for comp_name, scope_txt, rate_txt, base_sub, quote_sub in line_items:
        ws1.row_dimensions[cur_row].height = 22
        ws1.cell(row=cur_row, column=2, value=comp_name).font = REGULAR_FONT
        ws1.cell(row=cur_row, column=2).border = THIN_BORDER

        ws1.cell(row=cur_row, column=3, value=scope_txt).alignment = Alignment(horizontal="center")
        ws1.cell(row=cur_row, column=3).font = REGULAR_FONT
        ws1.cell(row=cur_row, column=3).border = THIN_BORDER

        ws1.cell(row=cur_row, column=4, value=rate_txt).alignment = Alignment(horizontal="center")
        ws1.cell(row=cur_row, column=4).font = REGULAR_FONT
        ws1.cell(row=cur_row, column=4).border = THIN_BORDER

        c_base = ws1.cell(row=cur_row, column=5, value=base_sub)
        c_base.number_format = "$#,##0.00"
        c_base.font = REGULAR_FONT
        c_base.border = THIN_BORDER

        c_quote = ws1.cell(row=cur_row, column=6, value=quote_sub)
        c_quote.number_format = "$#,##0.00"
        c_quote.font = BOLD_FONT
        c_quote.border = THIN_BORDER

        cur_row += 1

    # Base Subtotal Row
    ws1.row_dimensions[cur_row].height = 24
    ws1.cell(row=cur_row, column=2, value="Subtotal (Base Operational Cost)").font = BOLD_FONT
    ws1.cell(row=cur_row, column=2).border = THIN_BORDER
    ws1.cell(row=cur_row, column=3, value="").border = THIN_BORDER
    ws1.cell(row=cur_row, column=4, value="").border = THIN_BORDER

    c_bs = ws1.cell(row=cur_row, column=5, value=summary.get("total_base_cost", 0.0))
    c_bs.number_format = "$#,##0.00"
    c_bs.font = BOLD_FONT
    c_bs.border = THIN_BORDER

    c_bs_q = ws1.cell(row=cur_row, column=6, value=summary.get("total_base_cost", 0.0))
    c_bs_q.number_format = "$#,##0.00"
    c_bs_q.font = REGULAR_FONT
    c_bs_q.border = THIN_BORDER
    cur_row += 1

    # Profit Margin Row
    ws1.row_dimensions[cur_row].height = 24
    ws1.cell(row=cur_row, column=2, value=f"Profit Margin / Commercial Markup ({margin_pct:.0f}%)").font = BOLD_FONT
    ws1.cell(row=cur_row, column=2).border = THIN_BORDER
    ws1.cell(row=cur_row, column=3, value="").border = THIN_BORDER
    ws1.cell(row=cur_row, column=4, value="").border = THIN_BORDER

    ws1.cell(row=cur_row, column=5, value="—").alignment = Alignment(horizontal="center")
    ws1.cell(row=cur_row, column=5).border = THIN_BORDER

    c_pm = ws1.cell(row=cur_row, column=6, value=summary.get("total_profit_margin", 0.0))
    c_pm.number_format = "$#,##0.00"
    c_pm.font = BOLD_FONT
    c_pm.border = THIN_BORDER
    cur_row += 1

    # Final Total Proposal Row
    ws1.row_dimensions[cur_row].height = 30
    for col_i in range(2, 7):
        ws1.cell(row=cur_row, column=col_i).fill = TOTAL_ROW_FILL
        ws1.cell(row=cur_row, column=col_i).border = TOTAL_DOUBLE_BORDER

    ws1.cell(row=cur_row, column=2, value="TOTAL QUOTED INVESTMENT (USD)").font = Font(name="Calibri", size=11, bold=True, color="1A237E")
    ws1.cell(row=cur_row, column=3, value="")
    ws1.cell(row=cur_row, column=4, value="")

    c_tb = ws1.cell(row=cur_row, column=5, value=summary.get("total_base_cost", 0.0))
    c_tb.number_format = "$#,##0.00"
    c_tb.font = BOLD_FONT

    c_tf = ws1.cell(row=cur_row, column=6, value=summary.get("total_final_cost", 0.0))
    c_tf.number_format = "$#,##0.00"
    c_tf.font = Font(name="Calibri", size=13, bold=True, color="1B5E20")
    cur_row += 3

    # Section: Scope Classification Summary
    ws1.cell(row=cur_row, column=2, value="PAGE COMPLEXITY BREAKDOWN").font = SECTION_FONT
    cur_row += 1

    complexity_headers = ["Tier", "Description & Characteristic", "Assigned Pages", "Manual Effort / Page", "API Cost / Page"]
    ws1.row_dimensions[cur_row].height = 24
    for c_idx, h_text in enumerate(complexity_headers, start=2):
        cell = ws1.cell(row=cur_row, column=c_idx, value=h_text)
        cell.fill = ACCENT_BLUE_FILL
        cell.font = WHITE_BOLD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER

    cur_row += 1

    tier_rows = [
        ("Simple", "Static content only: headings, lists, decorative images, links <= 60", f"{summary.get('simple_pages', 0)} pages", "1.0 hr", "$0.15", SIMPLE_FILL, SIMPLE_FONT),
        ("Medium", "Standard forms, single video/audio, data tables, dropdown nav, links > 60", f"{summary.get('medium_pages', 0)} pages", "2.5 hrs", "$0.60", MEDIUM_FILL, MEDIUM_FONT),
        ("Complex", "Dynamic popups, live updating aria-live, custom widgets, multi-step forms, carousels", f"{summary.get('complex_pages', 0)} pages", "5.5 hrs", "$1.80", COMPLEX_FILL, COMPLEX_FONT),
    ]

    for t_name, t_desc, t_count, t_hrs, t_api, t_fill, t_font in tier_rows:
        ws1.row_dimensions[cur_row].height = 22
        c_t = ws1.cell(row=cur_row, column=2, value=t_name)
        c_t.fill = t_fill
        c_t.font = t_font
        c_t.alignment = Alignment(horizontal="center", vertical="center")
        c_t.border = THIN_BORDER

        c_d = ws1.cell(row=cur_row, column=3, value=t_desc)
        c_d.font = REGULAR_FONT
        c_d.border = THIN_BORDER

        c_c = ws1.cell(row=cur_row, column=4, value=t_count)
        c_c.font = BOLD_FONT
        c_c.alignment = Alignment(horizontal="center", vertical="center")
        c_c.border = THIN_BORDER

        c_h = ws1.cell(row=cur_row, column=5, value=t_hrs)
        c_h.font = REGULAR_FONT
        c_h.alignment = Alignment(horizontal="center", vertical="center")
        c_h.border = THIN_BORDER

        c_a = ws1.cell(row=cur_row, column=6, value=t_api)
        c_a.font = REGULAR_FONT
        c_a.alignment = Alignment(horizontal="center", vertical="center")
        c_a.border = THIN_BORDER
        cur_row += 1

    cur_row += 2

    # Deliverables & Commercial Notes
    ws1.cell(row=cur_row, column=2, value="SCOPE OF DELIVERABLES & COMMERCIAL TERMS").font = SECTION_FONT
    cur_row += 1

    notes = [
        "1. Full WCAG 2.2 Level A and AA audit covering all 55 Success Criteria.",
        "2. Manual verification using NVDA & JAWS screen readers with full keyboard tab-order review.",
        "3. Executive Audit Workbook containing Test Case Matrix, Defect Logs with Screen Reader Reproduction steps.",
        "4. Exact code remediation guidance and technical snippets for engineering teams.",
        "5. Retesting support for identified defects within 30 days of report delivery.",
    ]
    for n in notes:
        ws1.row_dimensions[cur_row].height = 18
        ws1.merge_cells(f'B{cur_row}:F{cur_row}')
        cell = ws1.cell(row=cur_row, column=2, value=n)
        cell.font = REGULAR_FONT
        cell.alignment = Alignment(vertical="center")
        cur_row += 1

    # ══════════════════════════════════════════════════════════════════════════
    # SHEET 2: PAGE LINE-ITEMS BREAKDOWN
    # ══════════════════════════════════════════════════════════════════════════
    sheet2_cols = [
        ("Sl. No.", 8),
        ("Page Title", 32),
        ("Target URL", 40),
        ("Complexity", 14),
        ("Manual Effort", 14),
        ("Manual Cost", 14),
        ("API Cost", 12),
        ("Platform Fee", 14),
        ("Base Cost", 14),
        ("Margin (30%)", 14),
        ("Final Quote", 16),
        ("Classification Rationale & Triggers", 55),
    ]

    ws2.row_dimensions[1].height = 30
    for idx, (h_title, col_w) in enumerate(sheet2_cols, start=1):
        ws2.column_dimensions[get_column_letter(idx)].width = col_w
        c = ws2.cell(row=1, column=idx, value=h_title)
        c.fill = PRIMARY_BLUE_FILL
        c.font = WHITE_BOLD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = HEADER_BORDER

    ws2.freeze_panes = "A2"

    p_row = 2
    for p_idx, p in enumerate(pages, start=1):
        ws2.row_dimensions[p_row].height = 24

        # Sl. No.
        ws2.cell(row=p_row, column=1, value=p_idx).alignment = Alignment(horizontal="center", vertical="center")
        ws2.cell(row=p_row, column=1).border = THIN_BORDER

        # Page Title
        ws2.cell(row=p_row, column=2, value=_safe_val(p.get("title", "Untitled"))).border = THIN_BORDER

        # URL (Hyperlinked)
        u_cell = ws2.cell(row=p_row, column=3, value=_safe_val(p.get("url", "")))
        u_cell.font = LINK_FONT
        u_cell.hyperlink = p.get("url", "")
        u_cell.border = THIN_BORDER

        # Complexity
        comp = p.get("complexity", "simple").lower()
        comp_cell = ws2.cell(row=p_row, column=4, value=comp.capitalize())
        comp_cell.alignment = Alignment(horizontal="center", vertical="center")
        comp_cell.border = THIN_BORDER
        if comp == "complex":
            comp_cell.fill = COMPLEX_FILL
            comp_cell.font = COMPLEX_FONT
        elif comp == "medium":
            comp_cell.fill = MEDIUM_FILL
            comp_cell.font = MEDIUM_FONT
        else:
            comp_cell.fill = SIMPLE_FILL
            comp_cell.font = SIMPLE_FONT

        # Manual Effort
        hrs_c = ws2.cell(row=p_row, column=5, value=p.get("manual_hours", 0.0))
        hrs_c.number_format = '0.0 "hrs"'
        hrs_c.alignment = Alignment(horizontal="center", vertical="center")
        hrs_c.border = THIN_BORDER

        # Manual Cost
        mc_c = ws2.cell(row=p_row, column=6, value=p.get("manual_cost", 0.0))
        mc_c.number_format = "$#,##0.00"
        mc_c.border = THIN_BORDER

        # API Cost
        api_c = ws2.cell(row=p_row, column=7, value=p.get("api_cost", 0.0))
        api_c.number_format = "$#,##0.00"
        api_c.border = THIN_BORDER

        # Platform Fee
        pf_c = ws2.cell(row=p_row, column=8, value=p.get("platform_cost", 0.0))
        pf_c.number_format = "$#,##0.00"
        pf_c.border = THIN_BORDER

        # Base Cost
        bc_c = ws2.cell(row=p_row, column=9, value=p.get("base_cost", 0.0))
        bc_c.number_format = "$#,##0.00"
        bc_c.border = THIN_BORDER

        # Margin
        pm_c = ws2.cell(row=p_row, column=10, value=p.get("profit_margin", 0.0))
        pm_c.number_format = "$#,##0.00"
        pm_c.border = THIN_BORDER

        # Final Quote
        fq_c = ws2.cell(row=p_row, column=11, value=p.get("final_cost", 0.0))
        fq_c.number_format = "$#,##0.00"
        fq_c.font = Font(name="Calibri", size=10, bold=True, color="1B5E20")
        fq_c.border = THIN_BORDER

        # Rationale & Triggers
        triggers = p.get("triggers", [])
        trig_str = "; ".join(triggers) if triggers else p.get("summary_rationale", "")
        rat_cell = ws2.cell(row=p_row, column=12, value=_safe_val(trig_str))
        rat_cell.font = REGULAR_FONT
        rat_cell.border = THIN_BORDER

        p_row += 1

    # Total Row at bottom
    ws2.row_dimensions[p_row].height = 26
    for c_i in range(1, 13):
        ws2.cell(row=p_row, column=c_i).fill = TOTAL_ROW_FILL
        ws2.cell(row=p_row, column=c_i).border = TOTAL_DOUBLE_BORDER

    ws2.cell(row=p_row, column=2, value="TOTAL ESTIMATION").font = BOLD_FONT

    tot_hrs = ws2.cell(row=p_row, column=5, value=f"=SUM(E2:E{p_row - 1})")
    tot_hrs.number_format = '0.0 "hrs"'
    tot_hrs.font = BOLD_FONT
    tot_hrs.alignment = Alignment(horizontal="center")

    tot_mc = ws2.cell(row=p_row, column=6, value=f"=SUM(F2:F{p_row - 1})")
    tot_mc.number_format = "$#,##0.00"
    tot_mc.font = BOLD_FONT

    tot_api = ws2.cell(row=p_row, column=7, value=f"=SUM(G2:G{p_row - 1})")
    tot_api.number_format = "$#,##0.00"
    tot_api.font = BOLD_FONT

    tot_pf = ws2.cell(row=p_row, column=8, value=f"=SUM(H2:H{p_row - 1})")
    tot_pf.number_format = "$#,##0.00"
    tot_pf.font = BOLD_FONT

    tot_bc = ws2.cell(row=p_row, column=9, value=f"=SUM(I2:I{p_row - 1})")
    tot_bc.number_format = "$#,##0.00"
    tot_bc.font = BOLD_FONT

    tot_pm = ws2.cell(row=p_row, column=10, value=f"=SUM(J2:J{p_row - 1})")
    tot_pm.number_format = "$#,##0.00"
    tot_pm.font = BOLD_FONT

    tot_fq = ws2.cell(row=p_row, column=11, value=f"=SUM(K2:K{p_row - 1})")
    tot_fq.number_format = "$#,##0.00"
    tot_fq.font = Font(name="Calibri", size=11, bold=True, color="1B5E20")

    # ══════════════════════════════════════════════════════════════════════════
    # SHEET 3: ELEMENT INVENTORY
    # ══════════════════════════════════════════════════════════════════════════
    inv_headers = [
        ("Sl. No.", 8),
        ("Target URL", 40),
        ("Complexity", 14),
        ("Total DOM Nodes", 16),
        ("Links", 10),
        ("Buttons", 10),
        ("Headings", 10),
        ("Paragraphs", 12),
        ("Lists", 10),
        ("Images / Icons", 14),
        ("Forms", 10),
        ("Inputs", 10),
        ("Tables", 10),
        ("Media / Embeds", 14),
        ("Specific Detected Triggers", 60),
    ]

    ws3.row_dimensions[1].height = 30
    for idx, (h_title, col_w) in enumerate(inv_headers, start=1):
        ws3.column_dimensions[get_column_letter(idx)].width = col_w
        c = ws3.cell(row=1, column=idx, value=h_title)
        c.fill = NAVY_HEADER_FILL
        c.font = WHITE_BOLD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = HEADER_BORDER

    ws3.freeze_panes = "A2"

    i_row = 2
    for p_idx, p in enumerate(pages, start=1):
        ws3.row_dimensions[i_row].height = 22
        counts = p.get("counts", {})

        ws3.cell(row=i_row, column=1, value=p_idx).alignment = Alignment(horizontal="center")
        ws3.cell(row=i_row, column=1).border = THIN_BORDER

        u_c = ws3.cell(row=i_row, column=2, value=_safe_val(p.get("url", "")))
        u_c.font = LINK_FONT
        u_c.hyperlink = p.get("url", "")
        u_c.border = THIN_BORDER

        comp = p.get("complexity", "simple").lower()
        comp_cell = ws3.cell(row=i_row, column=3, value=comp.capitalize())
        comp_cell.alignment = Alignment(horizontal="center")
        comp_cell.border = THIN_BORDER
        if comp == "complex":
            comp_cell.fill = COMPLEX_FILL
            comp_cell.font = COMPLEX_FONT
        elif comp == "medium":
            comp_cell.fill = MEDIUM_FILL
            comp_cell.font = MEDIUM_FONT
        else:
            comp_cell.fill = SIMPLE_FILL
            comp_cell.font = SIMPLE_FONT

        elem_vals = [
            counts.get("total_elements", 0),
            counts.get("links", 0),
            counts.get("buttons", 0),
            counts.get("headings", 0),
            counts.get("paragraphs", 0),
            counts.get("lists", 0),
            counts.get("images", 0),
            counts.get("forms", 0),
            counts.get("inputs", 0),
            counts.get("tables", 0),
            counts.get("media", 0),
        ]

        for offset, val in enumerate(elem_vals, start=4):
            c_val = ws3.cell(row=i_row, column=offset, value=val)
            c_val.alignment = Alignment(horizontal="center")
            c_val.border = THIN_BORDER

        trig_list = p.get("triggers", [])
        trig_cell = ws3.cell(row=i_row, column=15, value=_safe_val("; ".join(trig_list)))
        trig_cell.font = REGULAR_FONT
        trig_cell.border = THIN_BORDER

        i_row += 1

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream


def generate_quotation_pdf(data: Dict[str, Any]) -> io.BytesIO:
    """
    Generates an executive vector PDF Quotation Proposal document.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    normal = styles['Normal']

    header_style = ParagraphStyle(
        'DocHeader',
        parent=normal,
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0D1B2A')
    )
    sub_style = ParagraphStyle(
        'DocSub',
        parent=normal,
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#555555')
    )
    section_style = ParagraphStyle(
        'DocSection',
        parent=normal,
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#1A237E'),
        spaceAfter=6
    )
    cell_bold = ParagraphStyle(
        'CellBold',
        parent=normal,
        fontName='Helvetica-Bold',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#1A237E')
    )
    cell_normal = ParagraphStyle(
        'CellNormal',
        parent=normal,
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#222222')
    )
    cell_right_bold = ParagraphStyle(
        'CellRightBold',
        parent=cell_bold,
        alignment=2
    )

    story = []

    # Title & Metadata
    now = datetime.now()
    summary = data.get("summary", {})
    pages = data.get("pages", [])
    quote_ref = f"A11Y-QUOTE-{now.strftime('%Y%m%d')}-{abs(hash(str(summary))) % 10000:04d}"
    valid_until = (now + timedelta(days=30)).strftime("%B %d, %Y")

    story.append(Paragraph("A11ySense AI — Accessibility Audit Quotation", header_style))
    story.append(Paragraph(f"Commercial Proposal Ref: <b>{quote_ref}</b> | Issued: {now.strftime('%B %d, %Y')} | Valid: {valid_until}", sub_style))
    story.append(Spacer(1, 12))
    story.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor('#1A237E'), spaceAfter=15))

    # Target & Standards Table
    target_url = pages[0].get("url", "") if pages else ""
    meta_table_data = [
        [Paragraph("<b>Target Website:</b>", cell_normal), Paragraph(target_url, cell_bold), Paragraph("<b>Standards:</b>", cell_normal), Paragraph("WCAG 2.2 (Level A & AA), Section 508", cell_normal)],
        [Paragraph("<b>Total Scope:</b>", cell_normal), Paragraph(f"{summary.get('total_pages', len(pages))} Page(s) ({summary.get('simple_pages', 0)} Simple, {summary.get('medium_pages', 0)} Med, {summary.get('complex_pages', 0)} Complex)", cell_normal), Paragraph("<b>Engagement:</b>", cell_normal), Paragraph("Comprehensive Specialist Audit", cell_normal)],
    ]
    meta_table = Table(meta_table_data, colWidths=[1.2 * inch, 2.5 * inch, 1.2 * inch, 2.6 * inch])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('PADDING', (0, 0), (-1, -1), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 15))

    # Investment Summary
    story.append(Paragraph("Commercial Investment Summary", section_style))

    margin_pct = summary.get("profit_margin_pct", 30.0)
    margin_factor = 1.0 + (margin_pct / 100.0)

    manual_hrs = summary.get("total_manual_hours", 0.0)
    hourly_rate = summary.get("hourly_rate", 40.0)
    manual_base = summary.get("total_manual_cost", 0.0)
    manual_quoted = manual_base * margin_factor

    api_base = summary.get("total_api_cost", 0.0)
    api_quoted = api_base * margin_factor

    plat_base = summary.get("total_platform_cost", 0.0)
    plat_quoted = plat_base * margin_factor

    inv_table_data = [
        [Paragraph("<b>Service Component / Deliverable</b>", cell_bold), Paragraph("<b>Scope / Units</b>", cell_bold), Paragraph("<b>Unit Rate</b>", cell_bold), Paragraph("<b>Base Cost</b>", cell_bold), Paragraph("<b>Quoted Total (+30%)</b>", cell_right_bold)],
        [Paragraph("Manual Specialist Verification (NVDA, JAWS, Keyboard)", cell_normal), Paragraph(f"{manual_hrs:.1f} hrs", cell_normal), Paragraph(f"${hourly_rate:.2f}/hr", cell_normal), Paragraph(f"${manual_base:.2f}", cell_normal), Paragraph(f"${manual_quoted:.2f}", cell_right_bold)],
        [Paragraph("Automated WCAG 2.2 Rules & Model Engine", cell_normal), Paragraph(f"{summary.get('total_pages', len(pages))} pages", cell_normal), Paragraph("Tiered", cell_normal), Paragraph(f"${api_base:.2f}", cell_normal), Paragraph(f"${api_quoted:.2f}", cell_right_bold)],
        [Paragraph("Headless Browser Compute & Infrastructure", cell_normal), Paragraph(f"{summary.get('total_pages', len(pages))} pages", cell_normal), Paragraph(f"${summary.get('platform_fee_per_page', 10.0):.2f}/page", cell_normal), Paragraph(f"${plat_base:.2f}", cell_normal), Paragraph(f"${plat_quoted:.2f}", cell_right_bold)],
        [Paragraph("<b>Subtotal (Base Operational Cost)</b>", cell_bold), Paragraph("", cell_normal), Paragraph("", cell_normal), Paragraph(f"<b>${summary.get('total_base_cost', 0.0):.2f}</b>", cell_bold), Paragraph(f"${summary.get('total_base_cost', 0.0):.2f}", cell_right_bold)],
        [Paragraph(f"<b>Commercial Margin ({margin_pct:.0f}%)</b>", cell_bold), Paragraph("", cell_normal), Paragraph("", cell_normal), Paragraph("—", cell_normal), Paragraph(f"<b>${summary.get('total_profit_margin', 0.0):.2f}</b>", cell_right_bold)],
        [Paragraph("<b>TOTAL QUOTED INVESTMENT (USD)</b>", cell_bold), Paragraph("", cell_normal), Paragraph("", cell_normal), Paragraph(f"<b>${summary.get('total_base_cost', 0.0):.2f}</b>", cell_bold), Paragraph(f"<b>${summary.get('total_final_cost', 0.0):.2f}</b>", cell_right_bold)],
    ]

    inv_table = Table(inv_table_data, colWidths=[3.0 * inch, 1.1 * inch, 1.0 * inch, 1.1 * inch, 1.3 * inch])
    inv_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1A237E')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('BACKGROUND', (0, 4), (-1, 4), colors.HexColor('#F1F5F9')),
        ('BACKGROUND', (0, 6), (-1, 6), colors.HexColor('#E8F5E9')),
        ('PADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(inv_table)
    story.append(Spacer(1, 15))

    # Page Line Items
    story.append(Paragraph("Page-by-Page Audit Line Items", section_style))
    page_table_data = [
        [Paragraph("<b>#</b>", cell_bold), Paragraph("<b>Page Title & URL</b>", cell_bold), Paragraph("<b>Complexity</b>", cell_bold), Paragraph("<b>Hours</b>", cell_bold), Paragraph("<b>Final Cost</b>", cell_right_bold)]
    ]
    for idx, p in enumerate(pages, start=1):
        comp = p.get("complexity", "simple").upper()
        p_title = p.get("title", "Untitled")
        p_url = p.get("url", "")
        desc = f"<b>{p_title}</b><br/><font color='#555555'>{p_url}</font>"
        page_table_data.append([
            Paragraph(str(idx), cell_normal),
            Paragraph(desc, cell_normal),
            Paragraph(comp, cell_bold),
            Paragraph(f"{p.get('manual_hours', 0.0):.1f} hrs", cell_normal),
            Paragraph(f"<b>${p.get('final_cost', 0.0):.2f}</b>", cell_right_bold)
        ])

    page_table = Table(page_table_data, colWidths=[0.4 * inch, 4.3 * inch, 1.1 * inch, 0.8 * inch, 0.9 * inch])
    page_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0D1B2A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('PADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(page_table)
    story.append(Spacer(1, 15))

    # Terms & Acceptance
    story.append(Paragraph("Deliverables & Acceptance", section_style))
    terms_text = (
        "• Comprehensive WCAG 2.2 Level A/AA Test Case & Defect Matrix (Excel & Web Portal).<br/>"
        "• Screen reader reproduction steps (NVDA & JAWS) with exact code snippets for remediation.<br/>"
        "• 30-Day post-remediation verification included."
    )
    story.append(Paragraph(terms_text, cell_normal))
    story.append(Spacer(1, 20))

    # Signature Block
    sig_data = [
        [Paragraph("<b>Prepared by:</b> A11ySense AI Solutions Team", cell_normal), Paragraph("<b>Accepted by:</b> ___________________________", cell_normal)],
        [Paragraph("<b>Date:</b> " + now.strftime("%B %d, %Y"), cell_normal), Paragraph("<b>Signature / Date:</b> ___________________________", cell_normal)]
    ]
    sig_table = Table(sig_data, colWidths=[3.7 * inch, 3.8 * inch])
    sig_table.setStyle(TableStyle([
        ('PADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(sig_table)

    doc.build(story)
    buffer.seek(0)
    return buffer
