import io
import os
import re
import logging
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import urlparse

from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import _Cell
from docx.oxml.ns import qn

from common.constants.wcag import A11YSENSE_AUDIT_SCOPE, A11YSENSE_MANUAL_REVIEW_CRITERIA

logger = logging.getLogger(__name__)

TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "..", "templates", "accessibility_conformance_certificate_template.docx")

FIRM_NAME = os.getenv("CERT_FIRM_NAME", "WinVinaya")
AUDITOR_LINE = os.getenv("CERT_AUDITOR_NAME", "A11ySense Accessibility Audit Team, WinVinaya")
CONTACT_EMAIL = os.getenv("CERT_AUDITOR_EMAIL", "no-reply@winvinaya.com")

_MANUAL_CRITERIA_BY_CODE = {c["code"]: c for c in A11YSENSE_MANUAL_REVIEW_CRITERIA}
_AUTOMATED_CODES = {c["code"] for c in A11YSENSE_AUDIT_SCOPE}
_WCAG_CODE_RE = re.compile(r"(\d+\.\d+\.\d+)")


def _copy_run_format(new_run, ref_run) -> None:
    if ref_run is None:
        return
    new_run.font.name = ref_run.font.name
    new_run.font.size = ref_run.font.size
    new_run.font.bold = ref_run.font.bold
    new_run.font.italic = ref_run.font.italic
    if ref_run.font.color and ref_run.font.color.type is not None:
        new_run.font.color.rgb = ref_run.font.color.rgb


def _hyperlink_ref_run(paragraph: Paragraph):
    """python-docx's `paragraph.runs` skips runs wrapped in <w:hyperlink> (e.g. mailto:
    links), so grab one directly from the XML to use as a formatting reference."""
    for hyperlink in paragraph._p.findall(qn("w:hyperlink")):
        r = hyperlink.find(qn("w:r"))
        if r is not None:
            from docx.text.run import Run
            return Run(r, paragraph)
    return None


def _clear_paragraph_content(paragraph: Paragraph) -> None:
    """Removes all runs AND hyperlink-wrapped runs from a paragraph, since
    `run.text = ""` alone leaves hyperlink text (out of `paragraph.runs`) intact."""
    p = paragraph._p
    for child in list(p):
        if child.tag in (qn("w:r"), qn("w:hyperlink")):
            p.remove(child)


def _set_text_preserving_format(paragraph: Paragraph, text: str) -> None:
    """Replaces all content in a paragraph with a single run, copying formatting
    from whichever existing run had it (Word template runs often split one label
    across many runs due to autocorrect/spellcheck, and some end in a hyperlink)."""
    ref_run = next((r for r in paragraph.runs if r.text.strip()), None) or _hyperlink_ref_run(paragraph)
    _clear_paragraph_content(paragraph)
    run = paragraph.add_run(text)
    _copy_run_format(run, ref_run)


def _set_cell_text(cell: _Cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    ref_run = next((r for p in cell.paragraphs for r in p.runs if r.text.strip()), None)
    for extra_p in cell.paragraphs[1:]:
        extra_p._element.getparent().remove(extra_p._element)
    _clear_paragraph_content(paragraph)
    run = paragraph.add_run(text)
    _copy_run_format(run, ref_run)


def _append_to_paragraph(paragraph: Paragraph, suffix: str) -> None:
    ref_run = next((r for r in paragraph.runs if r.text.strip()), paragraph.runs[-1] if paragraph.runs else None)
    run = paragraph.add_run(suffix)
    _copy_run_format(run, ref_run)


def _build_criteria_index(testcases: List[Dict[str, Any]]) -> Dict[str, List[Dict[str, Any]]]:
    index: Dict[str, List[Dict[str, Any]]] = {}
    for tc in testcases:
        crit = tc.get("criteria", "N/A")
        if crit == "N/A" or not crit:
            continue
        code = crit.split(" ")[0]
        index.setdefault(code, []).append(tc)
    return index


def _compute_conformance(code: str, criteria_index: Dict[str, List[Dict[str, Any]]]) -> Tuple[str, str]:
    """Maps a WCAG success-criterion code to a VPAT-style conformance level and remark,
    derived from this audit's actual PASS/FAIL/NOT_APPLICABLE results."""
    manual = _MANUAL_CRITERIA_BY_CODE.get(code)
    if manual:
        if manual.get("group") == "A":
            remark = ("This criterion is not covered by A11ySense's automated evaluation and requires manual "
                       "testing; it may be automated in a future release.")
        else:
            remark = ("This criterion requires human judgement, real-device testing, or manual review of media "
                       "content, and was not covered by automated testing.")
        return "Not Evaluated", remark

    if code not in _AUTOMATED_CODES:
        # WCAG success criterion outside A11ySense's automated scope and not yet
        # tracked in the manual-review checklist (e.g. newer WCAG 2.2 additions).
        return "Not Evaluated", "This criterion is outside A11ySense's current automated evaluation scope and was not assessed."

    entries = criteria_index.get(code, [])
    if not entries:
        return "Not Applicable", "No elements matching this criterion were found on the audited pages."

    statuses = {e.get("status") for e in entries}
    fail_texts = list(dict.fromkeys(e.get("description") for e in entries if e.get("status") == "FAIL" and e.get("description")))
    pass_texts = list(dict.fromkeys(e.get("description") for e in entries if e.get("status") == "PASS" and e.get("description")))

    if "FAIL" in statuses and "PASS" in statuses:
        level = "Partially Supports"
    elif "FAIL" in statuses:
        level = "Does Not Support"
    elif "PASS" in statuses:
        level = "Supports"
    else:
        return "Not Applicable", "No elements matching this criterion were found on the audited pages."

    if level in ("Does Not Support", "Partially Supports"):
        remark = " ".join(fail_texts[:3]) or "Defects were identified against this criterion during automated testing."
    else:
        remark = " ".join(pass_texts[:3]) or "Tested and verified compliant during automated evaluation."
    return level, remark


def _fill_criteria_tables(doc: Document, criteria_index: Dict[str, List[Dict[str, Any]]]) -> None:
    """Overwrites the Conformance Level / Remarks cells in every criteria table with
    this audit's real results, and drops the stale page-overflow continuation rows
    that carried the template's original (now-replaced) remarks text."""
    for table in doc.tables[1:]:  # table[0] is the fixed Standards/Guideline table
        rows_to_delete = []
        for row in table.rows:
            cells = row.cells
            if len(cells) < 3:
                continue
            code_text = cells[0].text.strip()
            conf_text = cells[1].text.strip()

            if code_text == "Criteria" and conf_text == "Conformance Level":
                continue  # repeated header row

            if not code_text and not conf_text:
                rows_to_delete.append(row)  # page-overflow continuation row
                continue

            match = _WCAG_CODE_RE.search(code_text)
            if not match:
                continue
            code = match.group(1)
            conformance, remark = _compute_conformance(code, criteria_index)
            _set_cell_text(cells[1], conformance)
            _set_cell_text(cells[2], remark)

        for row in rows_to_delete:
            row._element.getparent().remove(row._element)


def _fill_cover_page(doc: Document, product_name: str, report_date: str, client_name: Optional[str]) -> None:
    for p in doc.paragraphs:
        text = p.text
        if text == "Name of the Firm":
            _set_text_preserving_format(p, FIRM_NAME)
        elif text.startswith("Name of Product:"):
            _append_to_paragraph(p, product_name)
        elif text.startswith("Report Date:"):
            _append_to_paragraph(p, report_date)
        elif text.startswith("Contact Information:"):
            _set_text_preserving_format(p, f"Contact Information:  Auditor: {AUDITOR_LINE}, {CONTACT_EMAIL}")
        elif text.strip().startswith("Client:"):
            if client_name:
                _append_to_paragraph(p, client_name)
        elif text.startswith("The accessibility evaluation was carried out"):
            _set_text_preserving_format(p, (
                "The accessibility evaluation was carried out using A11ySense, an automated accessibility "
                "auditing platform built on axe-core, WCAG 2.2 heuristics, and AI-assisted analysis. Automated "
                f"testing covers {len(A11YSENSE_AUDIT_SCOPE)} of the WCAG 2.1/2.2 Level A and AA success criteria; "
                "the remaining criteria require manual verification and are marked 'Not Evaluated' below. Findings "
                "are supplemented by AI-assisted review to validate accuracy."
            ))
        elif text.startswith("Testing was performed on the following"):
            _set_text_preserving_format(p, "Testing was performed using the following automated tooling:")
        elif text.startswith("Operating System:"):
            _set_text_preserving_format(p, "Testing Method: Automated headless browser scan (Chromium via Playwright)")
        elif text.startswith("Browser:"):
            _set_text_preserving_format(p, "Rule Engine: axe-core, extended with A11ySense WCAG heuristics")
        elif text.startswith("Screen Reader:"):
            _set_text_preserving_format(p, "Additional Checks: Simulated screen reader announcement analysis")


def _product_name_from_url(url: Optional[str]) -> str:
    if not url:
        return "N/A"
    try:
        parsed = urlparse(url)
        return parsed.netloc or url
    except Exception:
        return url


class CertificateRepository:
    def generate_certificate_docx(
        self,
        task_id: str,
        testcases: List[Dict[str, Any]],
        product_url: Optional[str] = None,
        client_name: Optional[str] = None,
    ) -> io.BytesIO:
        """Generates the Accessibility Conformance Certificate (.docx) for a completed
        audit, using the fixed WinVinaya/A11ySense template and this audit's real
        per-criterion results."""
        doc = Document(TEMPLATE_PATH)

        if not product_url and testcases:
            product_url = testcases[0].get("page_url")
        product_name = _product_name_from_url(product_url)
        report_date = datetime.now().strftime("%B %d, %Y")

        criteria_index = _build_criteria_index(testcases)

        _fill_cover_page(doc, product_name, report_date, client_name)
        _fill_criteria_tables(doc, criteria_index)

        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        return output


certificate_repo = CertificateRepository()
